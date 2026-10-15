#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
extrator_ofertas.py
===================
Lê a planilha de ofertas divulgada a cada período (aba "Turmas", em .xlsx ou .csv)
e gera um JSON limpo, pronto para o motor de planejamento de grade.

Uso básico:
    python extrator_ofertas.py "Planejado-2026-2 - Visualização.xlsx"

Opções:
    --saida      caminho do JSON de saída (padrão: ofertas_<periodo>.json)
    --periodo    período, ex.: 2026-2 (padrão: deduzido do nome do arquivo)
    --curso      prefixo do código do curso na planilha (padrão: BCC)
    --todas      não filtra por curso; exporta todas as turmas da planilha
    --curriculo  JSON do currículo, para conferir códigos e créditos (opcional)
    --aba        nome da aba, no caso de .xlsx (padrão: Turmas)

Dependências: apenas a biblioteca padrão do Python. Para ler .xlsx é preciso o openpyxl
(pip install openpyxl). Arquivos .csv não precisam de nada extra.

Como a planilha é organizada (e o que o script faz com isso)
------------------------------------------------------------
Cada turma ocupa um BLOCO de linhas, que começa na linha em que "Disciplina" E "Docente"
estão preenchidos. Dentro do bloco existem duas listas independentes, lado a lado:
  * TURMA: para quais cursos há vagas (colunas Código/Nome, Abrir?, Núm., Restrição);
  * AULAS: dias e horários de cada encontro (colunas Dia, Horário, Tipo, Sala, Lotação).

ATENÇÃO: uma linha com "Disciplina" preenchida e SEM docente não é uma turma nova. É a
mesma turma listada sob outro código (código antigo do currículo anterior, disciplina da
pós-graduação etc.), e os encontros da turma ficam espalhados entre essas linhas. Exemplo:
Segurança da Informação (ICP473) tem a aula de quarta na linha do ICP473 e a de sexta na
linha do ICP719 (Segurança em Redes), somando as 4 h semanais. O script junta tudo em um
único registro, guarda os outros códigos em "codigos_cruzados" e registra, em cada
encontro e em cada vaga, o código sob o qual a linha aparece ("listada_em").
"""

import argparse
import csv
import json
import re
import sys
from datetime import datetime
from pathlib import Path

# ----------------------------------------------------------------------------------------
# Posição de cada coluna na planilha (a planilha tem 15 colunas, de A até O)
# ----------------------------------------------------------------------------------------
POSICAO = {
    "docente": 0, "disciplina": 1, "periodo_tipo": 2, "horas": 3, "vagas": 4,
    "turma_curso": 5, "abrir": 6, "numero_siga": 7, "restricao": 8,
    "dia": 9, "horario": 10, "tipo_aula": 11, "sala": 12, "tipo_sala": 13, "lotacao": 14,
}
NUMERO_DE_COLUNAS = 15

# Textos esperados nas duas linhas de cabeçalho. Se a universidade mudar o layout da
# planilha, a comparação falha e o script para com uma mensagem clara, em vez de gerar
# um JSON errado em silêncio.
CABECALHO_LINHA_1 = {0: "docente", 1: "disciplina", 2: "período / tipo", 3: "horas",
                     4: "vagas siga", 5: "turma", 9: "aulas", 12: "salas"}
CABECALHO_LINHA_2 = {5: "código / nome", 6: "abrir", 7: "núm", 8: "restrição", 9: "dia",
                     10: "horário", 11: "tipo", 12: "sala", 13: "tipo", 14: "lotação"}

# Colunas (posição) cujo título na 2ª linha do cabeçalho traz um contador que muda a cada versão
COLUNAS_COM_CONTADOR = {6, 7}   # "Abrir? S:88 N:0" e "Núm. (87)"

DIAS_DA_SEMANA = ["SEG", "TER", "QUA", "QUI", "SEX", "SÁB", "DOM"]

# Como cada valor da coluna "Abrir?" é traduzido para uma situação legível
SITUACAO_POR_VALOR_DE_ABRIR = {
    "SIM": "confirmada",    # o instituto confirmou que a turma abre
    "NÃO": "cancelada",
    "NAO": "cancelada",
    "N/A": "externa",       # turma de outra unidade (ex.: Matemática), sem decisão do IC
    "": "pendente",         # sem informação na planilha
}

# Categoria da turma, a partir da coluna "Período / Tipo"
CATEGORIA_POR_TIPO = {"SVÇ": "servico", "ELET": "eletiva", "PPGI": "pos_graduacao",
                      "PROFCOMP": "pos_graduacao"}


class ErroDeLayout(Exception):
    """A planilha não tem o formato esperado."""


# ----------------------------------------------------------------------------------------
# 1) Leitura do arquivo
# ----------------------------------------------------------------------------------------
def normalizar_celula(valor):
    """Transforma qualquer célula em texto limpo (sem quebras de linha nem espaços extras)."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)  # o Excel guarda 5837 como 5837.0
    return re.sub(r"\s+", " ", str(valor)).strip()


def ler_linhas_da_planilha(caminho, nome_da_aba):
    """Devolve a planilha como lista de linhas, cada uma com exatamente 15 textos."""
    caminho = Path(caminho)
    sufixo = caminho.suffix.lower()

    if sufixo == ".csv":
        with open(caminho, newline="", encoding="utf-8-sig") as arquivo:
            linhas_brutas = list(csv.reader(arquivo))
    elif sufixo in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError:
            sys.exit("Para ler .xlsx instale o openpyxl:  pip install openpyxl")
        livro = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
        if nome_da_aba not in livro.sheetnames:
            raise ErroDeLayout(f'Aba "{nome_da_aba}" não encontrada. Abas disponíveis: {livro.sheetnames}')
        linhas_brutas = [list(linha) for linha in livro[nome_da_aba].iter_rows(values_only=True)]
    else:
        raise ErroDeLayout(f"Formato não suportado: {sufixo} (use .xlsx ou .csv)")

    linhas = []
    for linha in linhas_brutas:
        celulas = [normalizar_celula(c) for c in linha][:NUMERO_DE_COLUNAS]
        celulas += [""] * (NUMERO_DE_COLUNAS - len(celulas))  # completa linhas curtas
        linhas.append(celulas)
    return linhas


def localizar_cabecalho_e_validar(linhas):
    """Acha as duas linhas de cabeçalho, confere o layout e devolve (índice da 1ª linha de dados, aviso da planilha)."""
    indice_cabecalho = next((i for i, l in enumerate(linhas)
                             if l[0].casefold() == "docente" and l[1].casefold() == "disciplina"), None)
    if indice_cabecalho is None or indice_cabecalho + 1 >= len(linhas):
        raise ErroDeLayout('Não encontrei o cabeçalho ("Docente", "Disciplina", ...) na planilha.')

    divergencias = []
    for esperado, linha in ((CABECALHO_LINHA_1, linhas[indice_cabecalho]),
                            (CABECALHO_LINHA_2, linhas[indice_cabecalho + 1])):
        for coluna, texto_esperado in esperado.items():
            texto_encontrado = linha[coluna].casefold()
            if coluna in COLUNAS_COM_CONTADOR:
                correto = texto_encontrado.startswith(texto_esperado)  # título + contador variável
            else:
                correto = texto_encontrado == texto_esperado           # título exato
            if not correto:
                divergencias.append(f'coluna {coluna + 1}: esperado "{texto_esperado}", encontrado "{linha[coluna]}"')
    if divergencias:
        raise ErroDeLayout("O layout da planilha mudou:\n  " + "\n  ".join(divergencias))

    # A primeira linha costuma trazer um aviso (ex.: "Informações provisórias")
    aviso_da_planilha = next((l[0] for l in linhas[:indice_cabecalho] if l[0]), "")
    return indice_cabecalho + 2, aviso_da_planilha


# ----------------------------------------------------------------------------------------
# 2) Conversões de texto para valores utilizáveis
# ----------------------------------------------------------------------------------------
def converter_horario_em_minutos(texto):
    """
    Converte uma faixa de horário em (início, fim), em minutos desde a meia-noite.
        "08-10h"      -> (480, 600)
        "18h30-20h10" -> (1110, 1210)
    Devolve None se o texto não tiver um desses formatos (ex.: "? h").
    """
    padrao = re.fullmatch(r"(\d{1,2})(?:h(\d{2}))?-(\d{1,2})(?:h(\d{2}))?h?", texto.replace(" ", ""))
    if not padrao:
        return None
    hora_inicio, minuto_inicio, hora_fim, minuto_fim = padrao.groups()
    inicio = int(hora_inicio) * 60 + int(minuto_inicio or 0)
    fim = int(hora_fim) * 60 + int(minuto_fim or 0)
    return inicio, fim


def formatar_minutos(minutos):
    return f"{minutos // 60:02d}:{minutos % 60:02d}"


def converter_inteiro(texto):
    return int(texto) if re.fullmatch(r"\d+", texto) else None


def separar_codigo_e_nome(texto):
    """'ICP131 - Programação de Computadores I' -> ('ICP131', 'Programação de Computadores I')."""
    padrao = re.fullmatch(r"([A-Za-z0-9]+)\s*-\s*(.+)", texto)
    return (padrao.group(1).upper(), padrao.group(2).strip()) if padrao else None


def interpretar_periodo_tipo(texto):
    """
    '1 T1'  -> categoria 'periodo', período recomendado 1, turma 1
    'Elet'  -> categoria 'eletiva';  'SVÇ' -> 'servico';  'PPGI' -> 'pos_graduacao'
    """
    padrao = re.fullmatch(r"(\d+)\s*T(\d+)", texto)
    if padrao:
        return "periodo", int(padrao.group(1)), int(padrao.group(2))
    return CATEGORIA_POR_TIPO.get(texto.upper(), "outra"), None, None


# ----------------------------------------------------------------------------------------
# 3) Agrupamento das linhas em blocos e montagem de cada turma
# ----------------------------------------------------------------------------------------
def agrupar_em_blocos(linhas, indice_primeira_linha_de_dados):
    """
    Cada bloco (uma turma real) começa numa linha com 'Disciplina' E 'Docente' preenchidos.
    Uma linha com 'Disciplina' mas sem docente é a mesma turma listada sob outro código e
    continua o bloco atual (ver o texto no topo do arquivo).
    """
    blocos = []
    for indice in range(indice_primeira_linha_de_dados, len(linhas)):
        linha = linhas[indice]
        numero_da_linha = indice + 1  # numeração 1-based, igual à do Excel
        tem_disciplina = bool(linha[POSICAO["disciplina"]])
        tem_docente = bool(linha[POSICAO["docente"]])
        comeca_nova_turma = tem_disciplina and (tem_docente or not blocos)
        if comeca_nova_turma:
            blocos.append([])
        if blocos:
            blocos[-1].append((numero_da_linha, linha))
    return blocos


def montar_turma(bloco, avisos):
    """Transforma um bloco de linhas em um registro de turma. Devolve None se o bloco for ilegível."""
    primeira_linha_numero, primeira_linha = bloco[0]
    codigo_e_nome = separar_codigo_e_nome(primeira_linha[POSICAO["disciplina"]])
    if codigo_e_nome is None:
        avisos.append({"tipo": "disciplina_ilegivel", "linha": primeira_linha_numero,
                       "detalhe": primeira_linha[POSICAO["disciplina"]]})
        return None
    codigo, nome = codigo_e_nome

    tipo_de_oferta = primeira_linha[POSICAO["periodo_tipo"]]
    categoria, periodo_recomendado, numero_da_turma = interpretar_periodo_tipo(tipo_de_oferta)

    docentes, numeros_siga, cursos, restricoes, aulas = [], [], [], [], []
    codigos_cruzados = []        # outros códigos sob os quais a mesma turma é listada
    valores_de_abrir = []
    horario_completo = True
    codigo_da_linha = codigo     # código sob o qual a linha atual está listada

    for numero_da_linha, linha in bloco:
        # --- disciplina desta linha: a primeira é a principal; as demais são códigos cruzados
        texto_da_disciplina = linha[POSICAO["disciplina"]]
        if texto_da_disciplina:
            disciplina_lida = separar_codigo_e_nome(texto_da_disciplina)
            if disciplina_lida:
                codigo_da_linha = disciplina_lida[0]
                ja_registrado = codigo_da_linha == codigo or any(
                    cruzado["codigo"] == codigo_da_linha for cruzado in codigos_cruzados)
                if not ja_registrado:
                    codigos_cruzados.append({
                        "codigo": codigo_da_linha, "nome": disciplina_lida[1],
                        "tipo_de_oferta": linha[POSICAO["periodo_tipo"]],
                        "numero_siga": linha[POSICAO["numero_siga"]] or None})

        # --- docentes (uma turma pode ter mais de um, em linhas diferentes)
        docente = linha[POSICAO["docente"]]
        if docente and docente not in docentes:
            docentes.append(docente)

        # --- lista TURMA: cursos com vagas, "Abrir?", número no SIGA e restrição
        curso_texto = linha[POSICAO["turma_curso"]]
        if curso_texto:
            curso = separar_codigo_e_nome(curso_texto)
            if curso:
                cursos.append({"curso": curso[0], "nome": curso[1],
                               "vagas": converter_inteiro(linha[POSICAO["vagas"]]),
                               "listada_em": codigo_da_linha})
        if linha[POSICAO["abrir"]]:
            valores_de_abrir.append(linha[POSICAO["abrir"]].upper())
        numero_siga = linha[POSICAO["numero_siga"]]
        if numero_siga and numero_siga not in numeros_siga:
            numeros_siga.append(numero_siga)
        restricao = linha[POSICAO["restricao"]]
        if restricao and restricao.upper() not in ("NÃO", "NAO") and restricao not in restricoes:
            restricoes.append(restricao)

        # --- lista AULAS: dia e horário de cada encontro
        dia, horario = linha[POSICAO["dia"]], linha[POSICAO["horario"]]
        if not dia and not horario:
            continue
        faixa = converter_horario_em_minutos(horario)
        dia_normalizado = dia.upper()
        if dia_normalizado == "SAB":
            dia_normalizado = "SÁB"
        if faixa is None or dia_normalizado not in DIAS_DA_SEMANA or faixa[1] <= faixa[0]:
            horario_completo = False
            avisos.append({"tipo": "horario_indefinido", "turma": codigo, "linha": numero_da_linha,
                           "detalhe": f'dia="{dia}" horário="{horario}"'})
            continue
        aulas.append({
            "dia": dia_normalizado, "dia_numero": DIAS_DA_SEMANA.index(dia_normalizado),
            "inicio_min": faixa[0], "fim_min": faixa[1],
            "inicio": formatar_minutos(faixa[0]), "fim": formatar_minutos(faixa[1]),
            "tipo": linha[POSICAO["tipo_aula"]] or None,          # T = teórica, P = prática
            "sala": linha[POSICAO["sala"]] or None,
            "tipo_sala": linha[POSICAO["tipo_sala"]] or None,
            "lotacao": converter_inteiro(linha[POSICAO["lotacao"]]),
            "listada_em": codigo_da_linha,
        })

    if not aulas:
        horario_completo = False  # turma sem nenhum horário utilizável

    valor_de_abrir = valores_de_abrir[0] if valores_de_abrir else ""
    return {
        "codigo": codigo,
        "nome": nome,
        "codigos_cruzados": codigos_cruzados,
        "docentes": docentes,
        "tipo_de_oferta": tipo_de_oferta,
        "categoria": categoria,
        "periodo_recomendado": periodo_recomendado,
        "numero_da_turma": numero_da_turma,
        "horas_semanais": converter_inteiro(primeira_linha[POSICAO["horas"]]),
        "numeros_siga": numeros_siga,
        "abrir": valor_de_abrir or None,
        "situacao": SITUACAO_POR_VALOR_DE_ABRIR.get(valor_de_abrir, "pendente"),
        "restrita_aos_cursos": restricoes,
        "vagas_por_curso": cursos,
        "aulas": aulas,
        "horario_completo": horario_completo,
        "linha_na_planilha": primeira_linha_numero,
    }


# ----------------------------------------------------------------------------------------
# 4) Filtro por curso
# ----------------------------------------------------------------------------------------
def turma_e_aberta_ao_curso(turma, prefixo_do_curso):
    """A turma tem vagas para o curso e não é restrita a outros cursos?"""
    tem_vaga_para_o_curso = any(c["curso"].startswith(prefixo_do_curso) for c in turma["vagas_por_curso"])
    if not tem_vaga_para_o_curso:
        return False
    restricoes = turma["restrita_aos_cursos"]
    return not restricoes or any(r.upper().startswith(prefixo_do_curso) for r in restricoes)


# ----------------------------------------------------------------------------------------
# 5) Conferência com o currículo (opcional)
# ----------------------------------------------------------------------------------------
def conferir_com_curriculo(turmas, caminho_do_curriculo, avisos):
    """Marca o código equivalente no currículo e avisa sobre códigos desconhecidos ou créditos diferentes."""
    with open(caminho_do_curriculo, encoding="utf-8") as arquivo:
        curriculo = json.load(arquivo)
    disciplinas = curriculo["disciplinas"]
    apelidos = curriculo.get("alias_codigo", {})

    for turma in turmas:
        codigo = turma["codigo"]
        # A matéria do currículo pode estar no código principal ou em um código cruzado
        candidatos = [codigo] + [cruzado["codigo"] for cruzado in turma["codigos_cruzados"]]
        turma["codigo_no_curriculo"] = None
        for candidato in candidatos:
            if candidato in disciplinas:
                turma["codigo_no_curriculo"] = candidato
                break
            if candidato in apelidos:  # código antigo, equivalente 1 para 1
                turma["codigo_no_curriculo"] = apelidos[candidato]
                break

        if turma["codigo_no_curriculo"] is None:
            avisos.append({"tipo": "codigo_fora_do_curriculo", "turma": turma["id"],
                           "linha": turma["linha_na_planilha"], "detalhe": f"{codigo} - {turma['nome']}"})
            continue
        if turma["codigo_no_curriculo"] != codigo:
            avisos.append({"tipo": "codigo_antigo", "turma": turma["id"], "linha": turma["linha_na_planilha"],
                           "detalhe": f"{codigo} equivale a {turma['codigo_no_curriculo']} no currículo"})

        disciplina = disciplinas[turma["codigo_no_curriculo"]]
        if disciplina.get("horas_extensao", 0) > 0:
            continue  # atividades de extensão não têm carga semanal comparável
        horas_esperadas = disciplina["carga_horaria"] / 15  # semestre de 15 semanas
        if turma["horas_semanais"] is not None and abs(turma["horas_semanais"] - horas_esperadas) > 0.01:
            avisos.append({"tipo": "horas_diferentes_do_curriculo", "turma": turma["id"],
                           "linha": turma["linha_na_planilha"],
                           "detalhe": f"planilha: {turma['horas_semanais']} h/semana; "
                                      f"currículo: {disciplina['carga_horaria']} h ({horas_esperadas:g} h/semana)"})


def conferir_carga_semanal_das_aulas(turmas, avisos):
    """
    Compara as horas semanais declaradas com a soma dos encontros listados. Se os encontros
    somam menos, é provável que falte um horário na planilha (e o motor veria a turma mais
    'livre' do que ela realmente é).
    """
    for turma in turmas:
        if turma["horas_semanais"] is None or not turma["aulas"]:
            continue
        horas_nas_aulas = sum(a["fim_min"] - a["inicio_min"] for a in turma["aulas"]) / 60
        if abs(horas_nas_aulas - turma["horas_semanais"]) > 0.01:
            avisos.append({"tipo": "aulas_nao_somam_as_horas", "turma": turma["id"],
                           "linha": turma["linha_na_planilha"],
                           "detalhe": f'{turma["codigo"]}: declara {turma["horas_semanais"]} h/semana, '
                                      f'mas os encontros listados somam {horas_nas_aulas:g} h'})


# ----------------------------------------------------------------------------------------
# 6) Programa principal
# ----------------------------------------------------------------------------------------
def deduzir_periodo(nome_do_arquivo):
    padrao = re.search(r"(20\d{2})[-_.]?([12])", nome_do_arquivo)
    return f"{padrao.group(1)}-{padrao.group(2)}" if padrao else None


def extrair(caminho, nome_da_aba, prefixo_do_curso, exportar_todas, periodo, caminho_do_curriculo):
    linhas = ler_linhas_da_planilha(caminho, nome_da_aba)
    indice_dados, aviso_da_planilha = localizar_cabecalho_e_validar(linhas)

    avisos = []
    turmas_da_planilha = [t for t in (montar_turma(b, avisos) for b in agrupar_em_blocos(linhas, indice_dados)) if t]

    ignoradas = {"sem_vaga_para_o_curso": 0, "restrita_a_outro_curso": 0}
    turmas = []
    for turma in turmas_da_planilha:
        if not exportar_todas:
            if not any(c["curso"].startswith(prefixo_do_curso) for c in turma["vagas_por_curso"]):
                ignoradas["sem_vaga_para_o_curso"] += 1
                continue
            if not turma_e_aberta_ao_curso(turma, prefixo_do_curso):
                ignoradas["restrita_a_outro_curso"] += 1
                continue
        turmas.append(turma)

    # Identificador único: código + ordem de aparição daquele código na planilha
    contagem_por_codigo = {}
    for turma in turmas:
        contagem_por_codigo[turma["codigo"]] = contagem_por_codigo.get(turma["codigo"], 0) + 1
        turma["id"] = f'{turma["codigo"]}-{contagem_por_codigo[turma["codigo"]]}'

    # Os avisos de horário foram criados antes do filtro; mantém só os das turmas exportadas
    linhas_exportadas = {t["linha_na_planilha"] for t in turmas}
    codigos_exportados = {t["codigo"] for t in turmas}
    avisos = [a for a in avisos if a["tipo"] != "horario_indefinido" or a["turma"] in codigos_exportados]
    for turma in turmas:
        if turma["horario_completo"] is False and not turma["aulas"]:
            avisos.append({"tipo": "turma_sem_horario", "turma": turma["id"],
                           "linha": turma["linha_na_planilha"], "detalhe": f'{turma["codigo"]} - {turma["nome"]}'})

    conferir_carga_semanal_das_aulas(turmas, avisos)
    if caminho_do_curriculo:
        conferir_com_curriculo(turmas, caminho_do_curriculo, avisos)

    saida = {
        "_meta": {
            "periodo": periodo,
            "curso_filtrado": None if exportar_todas else prefixo_do_curso,
            "arquivo_de_origem": Path(caminho).name,
            "gerado_em": datetime.now().isoformat(timespec="seconds"),
            "aviso_da_planilha": aviso_da_planilha,
            "turmas_na_planilha": len(turmas_da_planilha),
            "turmas_exportadas": len(turmas),
            "turmas_ignoradas": ignoradas,
            "turmas_com_codigos_cruzados": sum(1 for t in turmas if t["codigos_cruzados"]),
            "legenda": {
                "situacao": "confirmada = 'Abrir?' é SIM; externa = N/A (turma de outra unidade); pendente = sem informação; cancelada = NÃO",
                "horario_completo": "false quando falta horário ou há horário ilegível; o motor deve tratar com cuidado ou descartar",
                "aulas": "inicio_min e fim_min em minutos desde a meia-noite; dois encontros colidem se estão no mesmo dia e inicio_a < fim_b e inicio_b < fim_a",
                "codigos_cruzados": "outros códigos sob os quais a mesma turma é listada (código antigo, pós-graduação); os encontros da turma podem estar em qualquer uma das linhas",
                "listada_em": "código da linha da planilha em que o encontro ou a vaga aparece",
                "horas_semanais": "coluna 'Horas' da planilha: horas de aula por semana (ex.: 4 para 60 h no semestre)",
            },
        },
        "turmas": turmas,
        "avisos": avisos,
    }
    return saida


def imprimir_resumo(saida):
    meta, turmas = saida["_meta"], saida["turmas"]
    print(f'Período: {meta["periodo"]} | arquivo: {meta["arquivo_de_origem"]}')
    if meta["aviso_da_planilha"]:
        print(f'Aviso da planilha: {meta["aviso_da_planilha"]}')
    print(f'Turmas na planilha: {meta["turmas_na_planilha"]} | exportadas: {meta["turmas_exportadas"]} '
          f'| ignoradas: {meta["turmas_ignoradas"]}')
    contagem = {}
    for turma in turmas:
        contagem[turma["situacao"]] = contagem.get(turma["situacao"], 0) + 1
    print("Situação das turmas exportadas:", contagem)
    sem_horario = sum(1 for t in turmas if not t["horario_completo"])
    print(f"Turmas com horário incompleto: {sem_horario}")
    por_tipo = {}
    for aviso in saida["avisos"]:
        por_tipo[aviso["tipo"]] = por_tipo.get(aviso["tipo"], 0) + 1
    print("Avisos:", por_tipo or "nenhum")


def main():
    analisador = argparse.ArgumentParser(description="Extrai as turmas ofertadas da planilha do período.")
    analisador.add_argument("arquivo", help="planilha de ofertas (.xlsx ou .csv da aba Turmas)")
    analisador.add_argument("--saida", help="JSON de saída (padrão: ofertas_<periodo>.json)")
    analisador.add_argument("--periodo", help="ex.: 2026-2 (padrão: deduzido do nome do arquivo)")
    analisador.add_argument("--curso", default="BCC", help="prefixo do curso na planilha (padrão: BCC)")
    analisador.add_argument("--todas", action="store_true", help="exportar todas as turmas, sem filtrar por curso")
    analisador.add_argument("--curriculo", help="JSON do currículo, para conferir códigos e créditos")
    analisador.add_argument("--aba", default="Turmas", help="nome da aba (somente .xlsx)")
    argumentos = analisador.parse_args()

    periodo = argumentos.periodo or deduzir_periodo(Path(argumentos.arquivo).name)
    if not periodo:
        sys.exit("Não consegui deduzir o período do nome do arquivo. Informe com --periodo 2026-2")

    try:
        saida = extrair(argumentos.arquivo, argumentos.aba, argumentos.curso.upper(), argumentos.todas,
                        periodo, argumentos.curriculo)
    except ErroDeLayout as erro:
        sys.exit(f"ERRO: {erro}")

    caminho_de_saida = Path(argumentos.saida or f"ofertas_{periodo}.json")
    with open(caminho_de_saida, "w", encoding="utf-8") as arquivo:
        json.dump(saida, arquivo, ensure_ascii=False, indent=2)
    imprimir_resumo(saida)
    print(f"JSON salvo em: {caminho_de_saida}")


if __name__ == "__main__":
    main()
